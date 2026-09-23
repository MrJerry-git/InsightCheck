"""T03 表格读取接线：CSV/TXT 由 T03 负责，Excel 按 H03 契约读取。

分工见 H03 说明（PR #19）：CSV 编码/分隔符探测与文件、任务管理属于 T03；
Excel 只读值读取与行标准化属于 H03（``app.tabular_parsing``）。

H03 在 main 上尚未合并，因此本模块先提供同一约定的等价实现
（openpyxl 只读、首个非空行是表头、值统一转字符串、空表头显式报错），
并在 :func:`excel_reader_name` 中如实报告当前生效的实现，
避免把回退实现当成"已接入 H03"。
"""

from __future__ import annotations

import csv
import io
from io import BytesIO
from typing import Any

CSV_SUFFIXES = frozenset({".csv", ".txt"})
EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm"})
TABULAR_SUFFIXES = CSV_SUFFIXES | EXCEL_SUFFIXES


class TabularSourceError(ValueError):
    """表格内容无法按约定读取；消息说明具体原因。"""


def _h03_reader() -> Any | None:
    """返回 H03 的 Excel 读取实现；未合并时返回 None（不静默假装已接入）。"""

    try:
        from app.tabular_parsing import read_excel_rows
    except ImportError:
        return None
    return read_excel_rows


def excel_reader_name() -> str:
    return "h03:app.tabular_parsing.read_excel_rows" if _h03_reader() else "t03:builtin-excel"


def read_csv_matrix(raw: bytes) -> list[list[str]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise TabularSourceError("文件必须是 UTF-8 编码的 CSV 或 TXT") from exc
    return [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]


def _read_excel_matrix_openpyxl(content: bytes, sheet_index: int = 0) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对损坏文件抛多种异常，统一转显式错误
        raise TabularSourceError(f"无法读取 Excel 文件：{exc}") from exc
    try:
        if sheet_index >= len(workbook.sheetnames):
            raise TabularSourceError(
                f"工作表索引 {sheet_index} 超出范围（共 {len(workbook.sheetnames)} 个）"
            )
        sheet = workbook[workbook.sheetnames[sheet_index]]
        matrix: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value).strip() for value in row]
            if any(cells):
                matrix.append(cells)
        if not matrix:
            raise TabularSourceError("工作表为空")
        headers = matrix[0]
        if all(not header for header in headers):
            raise TabularSourceError("表头行为空")
        blank = [position + 1 for position, header in enumerate(headers) if not header]
        if blank:
            raise TabularSourceError(f"表头存在空列（列号 {blank}），请补齐或删除空列")
        return matrix
    finally:
        workbook.close()


def read_excel_matrix(content: bytes, sheet_index: int = 0) -> list[list[str]]:
    """优先使用 H03 的 Excel 读取实现，未合并时用等价内置实现。"""

    reader = _h03_reader()
    if reader is None:
        return _read_excel_matrix_openpyxl(content, sheet_index)
    rows = reader(content, sheet_index)
    if not rows:
        raise TabularSourceError("工作表为空")
    headers = list(rows[0].keys())
    return [headers, *[[str(row.get(header, "")) for header in headers] for row in rows]]


def read_tabular_matrix(raw: bytes, suffix: str, sheet_index: int = 0) -> list[list[str]]:
    """按后缀读取 CSV/TXT/XLSX，统一返回"行 × 单元格"字符串矩阵。"""

    normalized = suffix.lower()
    if normalized in CSV_SUFFIXES:
        return read_csv_matrix(raw)
    if normalized in EXCEL_SUFFIXES:
        return read_excel_matrix(raw, sheet_index)
    raise TabularSourceError(f"不支持的表格格式：{suffix}")
